"""
CRS Maslahat - Comprehensive Backend API Test Suite
Run: python tests/test_api.py
Requires: pip install requests python-docx
"""
import requests, json, sys, os, io, time
from datetime import datetime

BASE_URL = os.environ.get("API_BASE", "http://localhost:8001/api")

ADMIN_EMAIL    = "admin@bsimaslahat.co.id"
ADMIN_PASSWORD = "Admin@2026"
BU_EMAIL       = "bu@bsimaslahat.co.id"
BU_PASSWORD    = "Demo@2026"
LEGAL_EMAIL    = "legal@bsimaslahat.co.id"
LEGAL_PASSWORD = "Demo@2026"
MGMT_EMAIL     = "management@bsimaslahat.co.id"
MGMT_PASSWORD  = "Demo@2026"

PASS_SYM = "[PASS]"; FAIL_SYM = "[FAIL]"; SKIP_SYM = "[SKIP]"
results = {"passed": 0, "failed": 0, "skipped": 0}

def check(name, cond, detail=""):
    if cond:
        print(f"  {PASS_SYM} {name}"); results["passed"] += 1
    else:
        print(f"  {FAIL_SYM} {name}" + (f" -- {detail}" if detail else "")); results["failed"] += 1

def skip(name, reason):
    print(f"  {SKIP_SYM} {name} (skip: {reason})"); results["skipped"] += 1

def section(title):
    print(f"\n{'='*55}\n  {title}\n{'='*55}")

def login(email, password):
    r = requests.post(f"{BASE_URL}/auth/login", json={"email": email, "password": password}, timeout=10)
    return r.json().get("token") if r.status_code == 200 else None

def auth(token): return {"Authorization": f"Bearer {token}"}

def make_dummy_docx():
    try:
        from docx import Document
        doc = Document()
        doc.add_paragraph("PERJANJIAN KERJA SAMA")
        doc.add_paragraph("BSI Maslahat : No. 03/001/PKS/BSI MASLAHAT/2026")
        doc.add_paragraph("Tentang PENYALURAN DANA BEASISWA 2026")
        doc.add_paragraph("2. YAYASAN PENDIDIKAN NUSANTARA , yang beralamat di Jakarta")
        doc.add_paragraph("nilai kerja sama sebesar Rp 500.000.000")
        doc.add_paragraph("dimulai efektif sejak tanggal 01 Agustus 2026")
        buf = io.BytesIO(); doc.save(buf); buf.seek(0); return buf.read()
    except ImportError:
        return b"PK\x03\x04" + b"\x00"*100

# ---- Section 1: AUTH ----
section("1 - AUTH")
admin_token = login(ADMIN_EMAIL, ADMIN_PASSWORD)
check("Admin login", admin_token is not None, f"email={ADMIN_EMAIL}")
bu_token = login(BU_EMAIL, BU_PASSWORD)
check("Business Unit login", bu_token is not None)
legal_token = login(LEGAL_EMAIL, LEGAL_PASSWORD)
check("Legal Officer login", legal_token is not None)
mgmt_token = login(MGMT_EMAIL, MGMT_PASSWORD)
check("Management login", mgmt_token is not None)

r = requests.post(f"{BASE_URL}/auth/login", json={"email": BU_EMAIL, "password": "wrong"}, timeout=10)
check("Reject wrong password -> 401", r.status_code == 401)

if admin_token:
    r = requests.get(f"{BASE_URL}/auth/me", headers=auth(admin_token), timeout=10)
    check("GET /auth/me returns admin", r.status_code == 200 and r.json().get("role") == "admin")

# ---- Section 2: META & GUIDELINES ----
section("2 - META & GUIDELINES (REQ-02)")
r = requests.get(f"{BASE_URL}/meta/options", headers=auth(bu_token or ""), timeout=10)
check("GET /meta/options -> 200", r.status_code == 200)
if r.status_code == 200:
    itypes = r.json().get("institution_types", [])
    check("DKM in institution_types", "DKM" in itypes, f"got: {itypes}")
    check("Perkumpulan in institution_types", "Perkumpulan" in itypes)
    check("All 7 institution types present", len(itypes) == 7, f"got {len(itypes)}")

for itype in ["Yayasan", "Perusahaan (PT)", "Koperasi", "Instansi Pemerintah", "DKM", "Perkumpulan", "Perorangan"]:
    r = requests.get(f"{BASE_URL}/guidelines/{requests.utils.quote(itype)}", headers=auth(bu_token or ""), timeout=10)
    check(f"Guidelines '{itype}' has items", r.status_code == 200 and len(r.json().get("items", [])) > 0,
          f"status={r.status_code}, items={len(r.json().get('items', []))}")

# ---- Section 3: CONTRACTS CRUD ----
section("3 - CONTRACTS CRUD")
contract_id = None
payload = {
    "partner_name": "Yayasan Auto Test", "partner_pic_name": "PIC Auto",
    "partner_pic_phone": "+6281234567890", "partner_pic_email": "autotest@yayasan.id",
    "institution_type": "Yayasan", "agreement_title": "PKS Auto Test 2026",
    "contract_value": 250000000, "effective_date": "2026-08-27", "expiry_date": "2027-08-27",
    "owning_bu": "CRG", "bu_pic_name": "Ahmad Auto", "remarks": "Auto test"
}
if bu_token:
    r = requests.post(f"{BASE_URL}/contracts", json=payload, headers=auth(bu_token), timeout=10)
    check("POST /contracts -> 200", r.status_code == 200, f"status={r.status_code}, {r.text[:120]}")
    if r.status_code == 200:
        contract_id = r.json().get("id")
        check("Contract has id", bool(contract_id))
        check("Status is drafting", r.json().get("status") == "drafting")
        check("contract_id format PKS-XXXX", r.json().get("contract_id","").startswith("PKS-"))

    r = requests.get(f"{BASE_URL}/contracts", headers=auth(bu_token), timeout=10)
    check("GET /contracts (BU) -> 200", r.status_code == 200)

    if contract_id:
        r = requests.get(f"{BASE_URL}/contracts/{contract_id}", headers=auth(bu_token), timeout=10)
        check("GET /contracts/{id} -> 200", r.status_code == 200)
        check("derived_status in response", r.status_code == 200 and "derived_status" in r.json())

        r = requests.put(f"{BASE_URL}/contracts/{contract_id}", json={"partner_name": "Updated Auto"},
                         headers=auth(bu_token), timeout=10)
        check("PUT /contracts/{id} (BU in drafting) -> 200", r.status_code == 200)

    if mgmt_token and contract_id:
        r = requests.put(f"{BASE_URL}/contracts/{contract_id}", json={"partner_name": "Should Fail"},
                         headers=auth(mgmt_token), timeout=10)
        check("Management cannot PUT contract -> 403", r.status_code == 403)
else:
    skip("Contract CRUD", "no BU token")

# ---- Section 4: FILE UPLOAD & WATERMARK ----
section("4 - FILE UPLOAD & WATERMARK (REQ-01, REQ-03, REQ-04)")
if bu_token and contract_id:
    docx_bytes = make_dummy_docx()
    files = {"file": ("draft.docx", io.BytesIO(docx_bytes),
             "application/vnd.openxmlformats-officedocument.wordprocessingml.document")}
    r = requests.post(f"{BASE_URL}/contracts/{contract_id}/versions",
                      files=files, data={"version_label": "v1.0", "remarks": "draft awal"},
                      headers=auth(bu_token), timeout=30)
    check("BU upload .docx v1.0 -> 200", r.status_code == 200, f"{r.status_code}: {r.text[:100]}")
    if r.status_code == 200:
        ver = r.json()
        check("Version is v1.0", ver.get("version") == "v1.0")
        check("watermarked=False for BU .docx", ver.get("watermarked") == False)

    # Submit + legal pickup so legal can upload
    requests.patch(f"{BASE_URL}/contracts/{contract_id}/status",
                   json={"status": "submitted_for_review"}, headers=auth(bu_token), timeout=10)
    if legal_token:
        r_pickup = requests.patch(f"{BASE_URL}/contracts/{contract_id}/status",
                                  json={"status": "under_legal_review"}, headers=auth(legal_token), timeout=10)
        check("Legal picks up -> under_legal_review", r_pickup.status_code == 200)

        # Legal upload PDF -> should be watermarked
        fake_pdf = b"%PDF-1.4 test content"
        files_pdf = {"file": ("review.pdf", io.BytesIO(fake_pdf), "application/pdf")}
        r = requests.post(f"{BASE_URL}/contracts/{contract_id}/versions",
                          files=files_pdf, data={"remarks": "coretan legal"},
                          headers=auth(legal_token), timeout=30)
        check("Legal upload PDF -> 200", r.status_code == 200, f"{r.status_code}: {r.text[:100]}")
        if r.status_code == 200:
            check("watermarked=True for Legal PDF (REQ-03)", r.json().get("watermarked") == True)

        # Legal upload JPG -> also watermarked
        fake_jpg = b"\xff\xd8\xff\xe0\x00\x10JFIF\x00\x01\x01\x00" + b"\x00"*50
        files_jpg = {"file": ("coretan.jpg", io.BytesIO(fake_jpg), "image/jpeg")}
        r = requests.post(f"{BASE_URL}/contracts/{contract_id}/versions",
                          files=files_jpg, data={"remarks": "jpg coretan"},
                          headers=auth(legal_token), timeout=30)
        check("Legal upload JPG accepted (REQ-03) -> not 403", r.status_code != 403,
              f"status={r.status_code}")
else:
    skip("File upload", "missing token or contract_id")

# ---- Section 5: STATUS TRANSITIONS ----
section("5 - STATUS TRANSITIONS & HARD-STOP (REQ-04)")
nc_id = None
if bu_token:
    p2 = {**payload, "partner_name": "Yayasan Status Flow", "agreement_title": "PKS Status Test",
          "partner_pic_email": "statustest@test.id"}
    r = requests.post(f"{BASE_URL}/contracts", json=p2, headers=auth(bu_token), timeout=10)
    if r.status_code == 200: nc_id = r.json()["id"]

    if nc_id:
        r = requests.patch(f"{BASE_URL}/contracts/{nc_id}/status",
                           json={"status": "submitted_for_review"}, headers=auth(bu_token), timeout=10)
        check("BU: drafting -> submitted_for_review", r.status_code == 200)

        r = requests.patch(f"{BASE_URL}/contracts/{nc_id}/status",
                           json={"status": "drafting"}, headers=auth(bu_token), timeout=10)
        check("BU: recall back to drafting", r.status_code == 200)

        r = requests.patch(f"{BASE_URL}/contracts/{nc_id}/status",
                           json={"status": "submitted_for_review"}, headers=auth(bu_token), timeout=10)

        if legal_token:
            r = requests.patch(f"{BASE_URL}/contracts/{nc_id}/status",
                               json={"status": "revision_required"}, headers=auth(legal_token), timeout=10)
            check("Legal: reject at intake -> revision_required", r.status_code == 200)

            r = requests.patch(f"{BASE_URL}/contracts/{nc_id}/status",
                               json={"status": "submitted_for_review"}, headers=auth(bu_token), timeout=10)
            r = requests.patch(f"{BASE_URL}/contracts/{nc_id}/status",
                               json={"status": "under_legal_review"}, headers=auth(legal_token), timeout=10)
            r = requests.patch(f"{BASE_URL}/contracts/{nc_id}/status",
                               json={"status": "ready_for_signature"}, headers=auth(legal_token), timeout=10)
            check("Legal: approve -> ready_for_signature", r.status_code == 200)

            # Hard-stop: contract_value locked
            r = requests.put(f"{BASE_URL}/contracts/{nc_id}", json={"contract_value": 99999},
                             headers=auth(bu_token), timeout=10)
            check("Locked field blocked in ready_for_signature -> 403", r.status_code == 403)

            r = requests.patch(f"{BASE_URL}/contracts/{nc_id}/status",
                               json={"status": "pending_final_verification"}, headers=auth(bu_token), timeout=10)
            check("BU: ready_for_signature -> pending_final_verification", r.status_code == 200)

            r = requests.patch(f"{BASE_URL}/contracts/{nc_id}/status",
                               json={"status": "signed_active"}, headers=auth(legal_token), timeout=10)
            check("Legal: final verify -> signed_active", r.status_code == 200)

            r = requests.patch(f"{BASE_URL}/contracts/{nc_id}/status",
                               json={"status": "drafting"}, headers=auth(bu_token), timeout=10)
            check("BU cannot move signed_active back -> 403", r.status_code == 403)

            if mgmt_token:
                r = requests.patch(f"{BASE_URL}/contracts/{nc_id}/status",
                                   json={"status": "drafting"}, headers=auth(mgmt_token), timeout=10)
                check("Management cannot change status -> 403", r.status_code == 403)

# ---- Section 6: AUTO-FILL EXTRACT ----
section("6 - AUTO-FILL EXTRACT DOCX (REQ-01)")
if bu_token:
    docx_bytes = make_dummy_docx()
    files = {"file": ("test.docx", io.BytesIO(docx_bytes),
             "application/vnd.openxmlformats-officedocument.wordprocessingml.document")}
    r = requests.post(f"{BASE_URL}/contracts/extract-docx", files=files,
                      headers=auth(bu_token), timeout=30)
    check("POST /contracts/extract-docx -> 200 or 400", r.status_code in (200, 400),
          f"{r.status_code}: {r.text[:80]}")
    if r.status_code == 200:
        d = r.json()
        check("agreement_title extracted", bool(d.get("agreement_title")))
        check("partner_name extracted", bool(d.get("partner_name")))
else:
    skip("Extract docx", "no bu_token")

# ---- Section 7: DASHBOARD & NOTIFICATIONS ----
section("7 - DASHBOARD, ANALYTICS & NOTIFICATIONS")
if admin_token:
    r = requests.get(f"{BASE_URL}/dashboard/stats", headers=auth(admin_token), timeout=10)
    check("GET /dashboard/stats -> 200", r.status_code == 200)
    if r.status_code == 200:
        s = r.json()
        check("Stats has total_active", "total_active" in s)
        check("Stats has expiring_soon", "expiring_soon" in s)

    r = requests.get(f"{BASE_URL}/dashboard/analytics", headers=auth(admin_token), timeout=10)
    check("GET /dashboard/analytics -> 200", r.status_code == 200)
    if r.status_code == 200:
        check("Analytics has by_bu", "by_bu" in r.json())
        check("Analytics has total_value", "total_value" in r.json())

    r = requests.get(f"{BASE_URL}/notifications", headers=auth(admin_token), timeout=10)
    check("GET /notifications -> 200", r.status_code == 200)
    if r.status_code == 200:
        check("Notifications has items", "items" in r.json())
        check("Notifications has unread count", "unread" in r.json())

    r = requests.post(f"{BASE_URL}/admin/run-expiry-reminders", headers=auth(admin_token), timeout=30)
    check("POST /admin/run-expiry-reminders -> 200", r.status_code == 200)

# ---- Section 8: RBAC & SECURITY ----
section("8 - RBAC & SECURITY (REQ-06)")
if bu_token:
    r = requests.get(f"{BASE_URL}/users", headers=auth(bu_token), timeout=10)
    check("BU cannot access /users -> 403", r.status_code == 403)

    r = requests.get(f"{BASE_URL}/contracts", timeout=10)
    check("No token -> 401/403", r.status_code in (401, 403))

if legal_token:
    r = requests.post(f"{BASE_URL}/contracts", json=payload, headers=auth(legal_token), timeout=10)
    check("Legal cannot POST /contracts -> 403", r.status_code == 403)

if admin_token:
    r = requests.get(f"{BASE_URL}/admin/audit-log", headers=auth(admin_token), timeout=10)
    check("GET /admin/audit-log -> 200", r.status_code == 200)

    r = requests.get(f"{BASE_URL}/admin/security-policy", headers=auth(admin_token), timeout=10)
    check("GET /admin/security-policy -> 200", r.status_code == 200)
    if r.status_code == 200:
        sp = r.json()
        check("Policy has locked_statuses", "locked_statuses" in sp)
        check("Policy has allowed_transitions", "allowed_transitions" in sp)

    if contract_id:
        r = requests.delete(f"{BASE_URL}/contracts/{contract_id}", headers=auth(admin_token), timeout=10)
        check("DELETE /contracts blocked -> 405", r.status_code == 405)

# ---- Section 9: COMMENTS ----
section("9 - DUAL REVIEW COMMENTS")
if contract_id and bu_token:
    r = requests.post(f"{BASE_URL}/contracts/{contract_id}/comments",
                      json={"text": "Komentar auto test", "section": "draft"},
                      headers=auth(bu_token), timeout=10)
    check("POST /comments -> 200", r.status_code == 200, f"{r.status_code}: {r.text[:80]}")
    comment_id = r.json().get("id") if r.status_code == 200 else None

    # BU attempt to resolve comment must be blocked with 403
    if comment_id and bu_token:
        r = requests.post(f"{BASE_URL}/comments/{comment_id}/resolve",
                          headers=auth(bu_token), timeout=10)
        check("BU cannot resolve comment -> 403", r.status_code == 403)

    if comment_id and legal_token:
        r = requests.post(f"{BASE_URL}/comments/{comment_id}/resolve",
                          headers=auth(legal_token), timeout=10)
        check("Legal resolve comment -> 200", r.status_code == 200)

# ---- Section 10: NEW SECURITY & CONSISTENCY CHECKS ----
section("10 - SECURITY HARDENING & CONSISTENCY VALIDATION")
if admin_token:
    # 1. Security Headers
    r = requests.get(f"{BASE_URL}/", timeout=10)
    check("Response has X-Content-Type-Options: nosniff", r.headers.get("x-content-type-options") == "nosniff")
    check("Response has X-Frame-Options: SAMEORIGIN", r.headers.get("x-frame-options") == "SAMEORIGIN")

    # 2. Regex Search Sanitization
    r = requests.get(f"{BASE_URL}/contracts?q=[PKS].*+?^$", headers=auth(admin_token), timeout=10)
    check("Special regex search chars sanitized -> 200", r.status_code == 200)

    # 3. Export XLSX with submitted_for_review
    r = requests.get(f"{BASE_URL}/reports/portfolio.xlsx?status=submitted_for_review", headers=auth(admin_token), timeout=15)
    check("Export XLSX with submitted_for_review -> 200", r.status_code == 200)

    # 4. Export PDF with submitted_for_review
    r = requests.get(f"{BASE_URL}/reports/portfolio.pdf?status=submitted_for_review", headers=auth(admin_token), timeout=15)
    check("Export PDF with submitted_for_review -> 200", r.status_code == 200)

# ---- SUMMARY ----
section("SUMMARY")
total = results["passed"] + results["failed"] + results["skipped"]
runnable = total - results["skipped"]
pct = int(results["passed"] / max(runnable, 1) * 100)
print(f"\n  Total   : {total}")
print(f"  {PASS_SYM} Passed : {results['passed']}")
print(f"  {FAIL_SYM} Failed : {results['failed']}")
print(f"  {SKIP_SYM} Skipped: {results['skipped']}")
print(f"\n  Score   : {pct}% ({results['passed']}/{runnable} runnable tests)")
if results["failed"] > 0:
    print("\n  [WARN] Some tests FAILED")
    sys.exit(1)
else:
    print("\n  [OK] All runnable tests passed!")
    sys.exit(0)
